"""เคสแปลก ๆ จากรอบทดสอบเชิงรุก — เขียนตาม "พฤติกรรมที่ควรเป็น" ตอนเขียนครั้งแรก
ล้ม 10 ตัว (แก้ครบแล้ว) เก็บไว้กันบั๊กเดิมกลับมา

รัน:  python manage.py test booking.tests_probe_weird --noinput -v 2
"""
from datetime import timedelta, time as time_type

from django.contrib.auth import get_user_model
from django.utils import timezone

from booking.models import Booking, MaintenanceBlock, TermBooking
from booking.tests_booking_edge import BookingEdgeBase, aware, next_weekday

User = get_user_model()


class ApprovedBookingEditProbe(BookingEdgeBase):
    """ผู้จองแก้การจองที่อนุมัติแล้ว — ควรกลับไปรออนุมัติ ไม่ใช่อนุมัติค้างไว้"""

    def test_move_time_of_approved_booking_needs_reapproval(self):
        b = self.seed(aware(self.day, 10), aware(self.day, 12), status='approved', user=self.bot)
        r = self.client.patch(f'/api/bookings/{b.id}/', {
            'start_time': aware(self.day, 15).isoformat(),
            'end_time': aware(self.day, 20).isoformat(),
        }, format='json')
        b.refresh_from_db()
        self.assertEqual(r.status_code, 200, r.content)
        self.assertNotEqual(b.status, 'approved', 'ย้ายเวลาได้เองโดยยังคงสถานะอนุมัติ')

    def test_move_room_of_approved_booking_needs_reapproval(self):
        b = self.seed(aware(self.day, 10), aware(self.day, 12), status='approved', user=self.bot)
        r = self.client.patch(f'/api/bookings/{b.id}/', {'room': self.room_b.id}, format='json')
        b.refresh_from_db()
        self.assertEqual(r.status_code, 200, r.content)
        self.assertNotEqual(b.status, 'approved', 'ย้ายห้องได้เองโดยยังคงสถานะอนุมัติ')

    def test_reapproval_is_logged_and_notified(self):
        b = self.seed(aware(self.day, 10), aware(self.day, 12), status='approved', user=self.bot)
        self.client.patch(f'/api/bookings/{b.id}/', {
            'end_time': aware(self.day, 13).isoformat()}, format='json')
        b.refresh_from_db()
        self.assertEqual(b.status, 'pending')
        self.assertTrue(b.logs.filter(old_status='approved', new_status='pending').exists())
        self.assertTrue(b.notification_set.filter(title__contains='รออนุมัติใหม่').exists())

    def test_title_only_edit_keeps_approval(self):
        b = self.seed(aware(self.day, 10), aware(self.day, 12), status='approved', user=self.bot)
        r = self.client.patch(f'/api/bookings/{b.id}/', {'title': 'ชื่อใหม่'}, format='json')
        b.refresh_from_db()
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(b.status, 'approved')

    def test_admin_moving_booking_keeps_approval(self):
        b = self.seed(aware(self.day, 10), aware(self.day, 12), status='approved', user=self.bot)
        self.client.force_authenticate(user=self.admin)
        r = self.client.patch(f'/api/bookings/{b.id}/', {'room': self.room_b.id}, format='json')
        b.refresh_from_db()
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(b.status, 'approved')

    def test_grow_attendees_to_capacity_keeps_ok(self):
        b = self.seed(aware(self.day, 10), aware(self.day, 12), status='approved', user=self.bot)
        r = self.client.patch(f'/api/bookings/{b.id}/', {'attendees': 41}, format='json')
        self.assertEqual(r.status_code, 400)


class MaintenanceTimezoneProbe(BookingEdgeBase):
    """serializer ของช่วงปิดซ่อมเทียบเวลา UTC กับเวลาคาบเรียน (เวลาไทย)"""

    def setUp(self):
        super().setUp()
        self.client.force_authenticate(user=self.admin)

    def term(self, dow, sh, eh):
        return TermBooking.objects.create(
            user=self.user_b, room=self.room, subject_name='วิชาทดสอบ', attendees=10,
            day_of_week=dow, start_time=time_type(sh), end_time=time_type(eh),
            term_start=self.day - timedelta(days=30), term_end=self.day + timedelta(days=60),
            status='active')

    def maint(self, start, end):
        return self.client.post('/api/maintenance-blocks/', {
            'room': self.room.id, 'start_time': start.isoformat(), 'end_time': end.isoformat(),
            'reason': 'ซ่อมแอร์',
        }, format='json')

    def test_afternoon_maintenance_not_blocked_by_morning_class(self):
        # คาบเรียนอังคาร 09–10 น. / ปิดซ่อมอังคาร 16–17 น. (= 09–10 UTC) ไม่ทับกัน
        self.term(self.day.weekday(), 9, 10)
        r = self.maint(aware(self.day, 16), aware(self.day, 17))
        self.assertEqual(r.status_code, 201, r.content.decode())

    def test_early_morning_maintenance_not_blocked_by_previous_day_class(self):
        # คาบเรียนอังคาร 18–19 น. / ปิดซ่อมพุธ 01–02 น. (= อังคาร 18–19 UTC) ไม่ทับกัน
        self.term(self.day.weekday(), 18, 19)
        wed = self.day + timedelta(days=1)
        r = self.maint(aware(wed, 1), aware(wed, 2))
        self.assertEqual(r.status_code, 201, r.content.decode())

    def test_patch_maintenance_onto_existing_booking_rejected(self):
        self.seed(aware(self.day, 10), aware(self.day, 12), status='approved', user=self.bot)
        mb = MaintenanceBlock.objects.create(
            room=self.room, start_time=aware(self.day, 14), end_time=aware(self.day, 15),
            reason='x', status='scheduled', created_by=self.admin)
        r = self.client.patch(f'/api/maintenance-blocks/{mb.id}/', {
            'start_time': aware(self.day, 9).isoformat(),
            'end_time': aware(self.day, 13).isoformat(),
        }, format='json')
        self.assertEqual(r.status_code, 400, 'แก้ช่วงปิดซ่อมให้ทับการจองที่อนุมัติแล้วได้')

    def test_patch_maintenance_end_before_start_rejected(self):
        mb = MaintenanceBlock.objects.create(
            room=self.room, start_time=aware(self.day, 14), end_time=aware(self.day, 15),
            reason='x', status='scheduled', created_by=self.admin)
        r = self.client.patch(f'/api/maintenance-blocks/{mb.id}/', {
            'end_time': aware(self.day, 8).isoformat(),
        }, format='json')
        self.assertEqual(r.status_code, 400, 'ตั้งเวลาสิ้นสุดก่อนเวลาเริ่มได้')

    def test_patch_maintenance_to_free_slot_ok(self):
        mb = MaintenanceBlock.objects.create(
            room=self.room, start_time=aware(self.day, 14), end_time=aware(self.day, 15),
            reason='x', status='scheduled', created_by=self.admin)
        r = self.client.patch(f'/api/maintenance-blocks/{mb.id}/', {
            'end_time': aware(self.day, 16).isoformat()}, format='json')
        mb.refresh_from_db()
        self.assertEqual(r.status_code, 200, r.content.decode())
        self.assertEqual(mb.end_time, aware(self.day, 16))

    def test_patch_multiday_maintenance_over_later_class_rejected(self):
        # ปิดซ่อมอังคาร–พฤหัส ต้องชนคาบเรียนวันพุธ (เดิมตรวจแค่วันแรก)
        self.term((self.day + timedelta(days=1)).weekday(), 9, 10)
        mb = MaintenanceBlock.objects.create(
            room=self.room, start_time=aware(self.day, 14), end_time=aware(self.day, 15),
            reason='x', status='scheduled', created_by=self.admin)
        r = self.client.patch(f'/api/maintenance-blocks/{mb.id}/', {
            'end_time': aware(self.day + timedelta(days=2), 12).isoformat()}, format='json')
        self.assertEqual(r.status_code, 400, r.content.decode())

    def test_patch_closed_maintenance_rejected(self):
        mb = MaintenanceBlock.objects.create(
            room=self.room, start_time=aware(self.day, 14), end_time=aware(self.day, 15),
            reason='x', status='cancelled', created_by=self.admin)
        r = self.client.patch(f'/api/maintenance-blocks/{mb.id}/', {
            'end_time': aware(self.day, 16).isoformat()}, format='json')
        self.assertEqual(r.status_code, 400)


class TermBookingWeirdProbe(BookingEdgeBase):
    def test_term_entirely_in_the_past_rejected(self):
        start = timezone.localdate() - timedelta(days=200)
        r = self.client.post('/api/term-bookings/', {
            'room': self.room.id, 'subject_name': 'ย้อนอดีต', 'attendees': 10,
            'day_of_week': start.weekday(), 'start_time': '09:00', 'end_time': '10:00',
            'term_start': start.isoformat(), 'term_end': (start + timedelta(days=30)).isoformat(),
        }, format='json')
        self.assertEqual(r.status_code, 400, 'จองทั้งเทอมย้อนหลังทั้งช่วงได้')

    def test_patch_term_to_other_room_over_capacity_rejected(self):
        small = self.room_b
        small.capacity = 5
        small.save()
        tb = TermBooking.objects.create(
            user=self.bot, room=self.room, subject_name='วิชา', attendees=30,
            day_of_week=self.day.weekday(), start_time=time_type(9), end_time=time_type(10),
            term_start=self.day, term_end=self.day + timedelta(days=30), status='active')
        r = self.client.patch(f'/api/term-bookings/{tb.id}/', {'room': small.id}, format='json')
        self.assertEqual(r.status_code, 400)


class RegisterWeirdProbe(BookingEdgeBase):
    def setUp(self):
        super().setUp()
        self.client.force_authenticate(user=None)

    def reg(self, username, **extra):
        payload = {'username': username, 'first_name': 'ก', 'last_name': 'ข',
                   'email': f'probe{abs(hash(username))}@example.com',
                   'password': 'Probe12345', 'password2': 'Probe12345', 'role': 'student'}
        payload.update(extra)
        return self.client.post('/api/auth/register/', payload, format='json')

    def test_username_only_domain_rejected(self):
        r = self.reg('@ubu.ac.th')
        self.assertEqual(r.status_code, 400, 'สมัครได้ username ว่าง')
        self.assertFalse(User.objects.filter(username='').exists())

    def test_duplicate_email_rejected(self):
        r = self.reg('probe_dup_email', email='qa_test_bot@example.com')
        self.assertEqual(r.status_code, 400, 'สมัครด้วยอีเมลที่มีคนใช้แล้วได้')

    def test_forgot_password_does_not_reveal_ldap_account(self):
        ldap_user = User.objects.create(username='6699999999', email='6699999999@ubu.ac.th')
        ldap_user.set_unusable_password()
        ldap_user.save()
        a = self.client.post('/api/auth/forgot-password/', {'username': '6699999999'}, format='json')
        b = self.client.post('/api/auth/forgot-password/', {'username': 'no_such_user_xyz'}, format='json')
        self.assertEqual(a.json(), b.json(), 'คำตอบต่างกัน -> เดาได้ว่ารหัสนักศึกษานี้มีบัญชีในระบบ')


class BookingWeirdInputProbe(BookingEdgeBase):
    def test_booking_room_marked_disabled_rejected(self):
        self.room.status = 'disabled'
        self.room.save()
        r = self.book(aware(self.day, 10), aware(self.day, 11))
        self.assertEqual(r.status_code, 400, 'จองห้องที่สถานะ disabled ได้')

    def test_far_future_booking_rejected(self):
        d = timezone.localdate().replace(year=timezone.localdate().year + 50)
        r = self.book(aware(d, 10), aware(d, 11))
        self.assertEqual(r.status_code, 400, 'จองล่วงหน้า 50 ปีได้')

    def test_nul_byte_in_title(self):
        r = self.book(aware(self.day, 10), aware(self.day, 11), title='abc\x00def')
        self.assertIn(r.status_code, (201, 400), f'status {r.status_code}')

    def test_zero_length_booking_rejected(self):
        r = self.book(aware(self.day, 10), aware(self.day, 10))
        self.assertEqual(r.status_code, 400)

    def test_huge_attendees_number(self):
        r = self.book(aware(self.day, 10), aware(self.day, 11), attendees=10**20)
        self.assertEqual(r.status_code, 400, f'status {r.status_code}')

    def test_search_term_without_term_dates(self):
        r = self.client.post('/api/rooms/search/', {
            'attendees': 5, 'start_time': '09:00', 'end_time': '10:00',
            'booking_type': 'term', 'day_of_week': 1,
        }, format='json')
        self.assertLess(r.status_code, 500, r.content[:300])

    def test_search_garbage_params(self):
        for q in ('?date=2026-13-45&start_time=25:99', '?date=%00&start_time=%00',
                  '?attendees=-1', '?building=abc&floor=xyz'):
            r = self.client.get('/api/rooms/' + q)
            self.assertLess(r.status_code, 500, q)
        for path in ('/api/rooms/today-feed/?date=nope', '/api/rooms/status-feed/?date=nope',
                     '/api/rooms/1/availability/?date=nope', '/api/term-bookings/?room=abc',
                     '/api/maintenance/slots/?room=abc&date=nope'):
            r = self.client.get(path)
            self.assertLess(r.status_code, 500, path)

    def test_split_recommend_garbage(self):
        r = self.client.post('/api/term-bookings/split-recommend/', {
            'day_of_week': 1, 'start_time': '9', 'end_time': '10:00',
            'term_start': '2026-01-01', 'term_end': '2025-01-01', 'attendees': 'abc',
        }, format='json')
        self.assertLess(r.status_code, 500, r.content[:300])

    def test_split_book_garbage_dow(self):
        r = self.client.post('/api/term-bookings/split-book/', {
            'subject_name': 'x', 'day_of_week': 'abc', 'start_time': '09:00', 'end_time': '10:00',
            'first_room_id': 'x', 'first_term_start': '2026-01-01', 'first_term_end': '2026-02-01',
            'second_room_id': self.room_b.id, 'second_term_start': '2026-02-02',
            'second_term_end': '2026-03-01',
        }, format='json')
        self.assertLess(r.status_code, 500, r.content[:300])

    def test_profile_patch_email_of_other_user_rejected(self):
        r = self.client.patch('/api/auth/profile/', {'email': 'QA_USER_B@example.com'}, format='json')
        self.assertEqual(r.status_code, 400)

    def test_profile_patch_own_email_ok(self):
        r = self.client.patch('/api/auth/profile/', {'email': 'qa_test_bot@example.com'}, format='json')
        self.assertEqual(r.status_code, 200, r.content)

    def test_legacy_duplicate_email_can_still_save_profile(self):
        # บัญชีเก่าที่อีเมลซ้ำกันอยู่ก่อนมีกฎนี้ ต้องบันทึกโปรไฟล์ได้ถ้าไม่ได้เปลี่ยนอีเมล
        User.objects.filter(pk=self.user_b.pk).update(email='qa_test_bot@example.com')
        r = self.client.patch('/api/auth/profile/', {
            'first_name': 'ชื่อใหม่', 'email': 'qa_test_bot@example.com'}, format='json')
        self.assertEqual(r.status_code, 200, r.content)

    def test_booking_within_advance_limit_ok(self):
        d = timezone.localdate() + timedelta(days=300)
        r = self.book(aware(d, 10), aware(d, 11))
        self.assertEqual(r.status_code, 201, r.content)

    def test_term_started_in_past_still_running_ok(self):
        start = timezone.localdate() - timedelta(days=20)
        r = self.client.post('/api/term-bookings/', {
            'room': self.room.id, 'subject_name': 'เทอมที่เริ่มแล้ว', 'attendees': 10,
            'day_of_week': start.weekday(), 'start_time': '09:00', 'end_time': '10:00',
            'term_start': start.isoformat(), 'term_end': (start + timedelta(days=60)).isoformat(),
        }, format='json')
        self.assertEqual(r.status_code, 201, r.content)

    def test_profile_patch_student_id_of_other_user(self):
        r = self.client.patch('/api/auth/profile/', {'student_id': '6622222222'}, format='json')
        self.assertEqual(r.status_code, 400)


# ══════════════════════════════════════════════════════════════
# รอบ 2 — timezone, อีเมล, export, การค้นหาห้อง
# ══════════════════════════════════════════════════════════════
import io
from datetime import datetime
from unittest import mock

import openpyxl
from django.core import mail
from django.test import override_settings

from booking.models import DemandForecast, Room

LOCMEM = 'django.core.mail.backends.locmem.EmailBackend'


def thai_now(d, hh, mm=0):
    """timezone.now() ปลอม (คืน UTC เหมือนของจริง) ที่ตรงกับเวลาไทย d hh:mm"""
    return aware(d, hh, mm).astimezone(dt_timezone.utc)


from datetime import timezone as dt_timezone  # noqa: E402


@override_settings(EMAIL_BACKEND=LOCMEM, DEFAULT_FROM_EMAIL='noreply@example.com')
class EmailProbe(BookingEdgeBase):
    def test_checkin_reminder_shows_thai_time(self):
        from booking.scheduler import send_checkin_reminders
        start = timezone.now() + timedelta(minutes=15)
        Booking.objects.create(user=self.bot, room=self.room, title='ประชุม', attendees=5,
                               start_time=start, end_time=start + timedelta(hours=1),
                               status='approved')
        send_checkin_reminders()
        self.assertEqual(len(mail.outbox), 1)
        html = mail.outbox[0].alternatives[0][0]
        self.assertIn(timezone.localtime(start).strftime('%H:%M'), html)

    def test_reminder_escapes_title(self):
        from booking.scheduler import send_checkin_reminders
        start = timezone.now() + timedelta(minutes=15)
        Booking.objects.create(user=self.bot, room=self.room, title='<a href="http://evil">x</a>',
                               attendees=5, start_time=start, end_time=start + timedelta(hours=1),
                               status='approved')
        send_checkin_reminders()
        html = mail.outbox[0].alternatives[0][0]
        self.assertNotIn('<a href="http://evil">', html)

    def test_pending_email_escapes_title(self):
        from booking.signals import send_booking_pending_email
        b = self.seed(aware(self.day, 10), aware(self.day, 11), status='pending', user=self.bot)
        b.title = '<img src=x onerror=alert(1)>'
        send_booking_pending_email(b)
        html = mail.outbox[-1].alternatives[0][0]
        self.assertNotIn('<img src=x', html)


class ThaiDayProbe(BookingEdgeBase):
    def test_dashboard_counts_thai_today_after_midnight(self):
        # 01:30 เวลาไทย = 18:30 UTC ของเมื่อวาน
        today = self.day
        self.seed(aware(today, 9), aware(today, 10), status='approved')
        self.client.force_authenticate(user=self.admin)
        with mock.patch('django.utils.timezone.now', return_value=thai_now(today, 1, 30)):
            r = self.client.get('/api/dashboard/')
        self.assertEqual(r.json()['today_dynamic'], 1)

    def test_room_forecast_uses_thai_hour(self):
        self.client.force_authenticate(user=self.admin)
        DemandForecast.objects.create(room=self.room, forecast_date=self.day, hour=10,
                                      predicted_demand=0.9, demand_level='high')
        with mock.patch('django.utils.timezone.now', return_value=thai_now(self.day, 10, 20)):
            r = self.client.get('/api/rooms/')
        rows = r.json()['results'] if isinstance(r.json(), dict) else r.json()
        mine = next(x for x in rows if x['id'] == self.room.id)
        self.assertTrue(mine['forecast']['has_forecast'], mine['forecast'])

    def test_export_times_are_thai(self):
        b = self.seed(aware(self.day, 10), aware(self.day, 11), status='approved')
        self.client.force_authenticate(user=self.admin)
        r = self.client.get('/api/export/excel/?sheets=bookings')
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        ws = wb.active
        header = [c.value for c in ws[1]]
        start = ws.cell(row=2, column=header.index('start_time') + 1).value
        self.assertEqual(start.hour, 10, f'{start} (booking {b.id})')

    def test_availability_lists_pending_bookings(self):
        self.seed(aware(self.day, 10), aware(self.day, 11), status='pending')
        self.client.force_authenticate(user=self.admin)
        r = self.client.get(f'/api/rooms/{self.room.id}/availability/?date={self.day.isoformat()}')
        self.assertEqual(len(r.json()['dynamic_bookings']), 1)


class RoomSearchProbe(BookingEdgeBase):
    """การค้นหาจำกัดเฉพาะห้องใน AI_FORECAST_ROOM_IDS — สร้างห้องด้วย id นั้น"""

    def setUp(self):
        super().setUp()
        self.ai_room = Room.objects.create(id=443, building=self.building, name='2C05',
                                           floor=1, capacity=40, room_type='ห้องเรียน')

    def search(self, **params):
        base = {'attendees': 5, 'start_time': '13:00', 'end_time': '14:00'}
        base.update(params)
        r = self.client.post('/api/rooms/search/', base, format='json')
        self.assertEqual(r.status_code, 200, r.content[:300])
        return [x['id'] for x in r.json()]

    def test_room_under_repair_today_still_found_next_week(self):
        now = timezone.now()
        MaintenanceBlock.objects.create(room=self.ai_room, start_time=now - timedelta(hours=1),
                                        end_time=now + timedelta(hours=1), reason='x',
                                        status='active', created_by=self.admin)
        self.ai_room.status = 'maintenance'
        self.ai_room.save()
        ids = self.search(date=self.day.isoformat())
        self.assertIn(self.ai_room.id, ids)

    def test_disabled_room_not_found(self):
        self.ai_room.status = 'disabled'
        self.ai_room.save()
        self.assertNotIn(self.ai_room.id, self.search(date=self.day.isoformat()))

    def test_term_search_hides_room_with_dynamic_booking_on_that_weekday(self):
        self.seed(aware(self.day, 13), aware(self.day, 14), status='pending', room=self.ai_room)
        ids = self.search(booking_type='term', day_of_week=self.day.weekday(),
                          term_start=(self.day - timedelta(days=7)).isoformat(),
                          term_end=(self.day + timedelta(days=30)).isoformat())
        self.assertNotIn(self.ai_room.id, ids)

    def test_term_search_ignores_maintenance_on_other_weekday(self):
        other = self.day + timedelta(days=2)
        MaintenanceBlock.objects.create(room=self.ai_room, start_time=aware(other, 8),
                                        end_time=aware(other, 17), reason='x',
                                        status='scheduled', created_by=self.admin)
        ids = self.search(booking_type='term', day_of_week=self.day.weekday(),
                          term_start=self.day.isoformat(),
                          term_end=(self.day + timedelta(days=30)).isoformat())
        self.assertIn(self.ai_room.id, ids)

    def test_term_search_null_dates_ok(self):
        self.search(booking_type='term', day_of_week=1, term_start=None, term_end=None)
