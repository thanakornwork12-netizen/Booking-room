"""ทดสอบแบบจงใจหาช่องโหว่ — ข้อมูลแปลก, ปลอม token, ยึดบัญชี, export, ลิงก์อีเมล

รัน:  python manage.py test booking.tests_adversarial --noinput
"""
import io
from datetime import datetime, time as time_type, timedelta

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, TransactionTestCase
from django.utils import timezone
from rest_framework.test import APIClient
from rest_framework_simplejwt.tokens import RefreshToken

from booking.models import Booking, Building, Room

User = get_user_model()


def next_weekday(weekday, min_days=7):
    d = timezone.localdate() + timedelta(days=min_days)
    while d.weekday() != weekday:
        d += timedelta(days=1)
    return d


def aware(d, hh, mm=0):
    return timezone.make_aware(datetime.combine(d, time_type(hh, mm)))


class AdvBase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.building = Building.objects.create(name='อาคารเรียนรวม 2C', code='2C')
        cls.room = Room.objects.create(id=445, building=cls.building, name='2C09', floor=1,
                                       capacity=40, room_type='ห้องเรียน')

    def setUp(self):
        self.client = APIClient()
        self.client.raise_request_exception = False
        self.bot = User.objects.create_user(username='qa_test_bot', password='QaTest1234',
                                            role='student', email='qa_test_bot@example.com')
        self.admin = User.objects.create_user(username='qa_admin', password='QaAdmin1234',
                                              role='admin', email='qa_admin@example.com')
        self.client.force_authenticate(user=self.bot)
        self.tue = next_weekday(1)

    def book(self, **over):
        p = {'room': self.room.id, 'title': 'ประชุม', 'attendees': 10,
             'start_time': aware(self.tue, 13).isoformat(), 'end_time': aware(self.tue, 15).isoformat()}
        p.update(over)
        return self.client.post('/api/bookings/', p, format='json')

    def bearer(self, token):
        c = APIClient()
        c.raise_request_exception = False
        c.credentials(HTTP_AUTHORIZATION=f'Bearer {token}')
        return c


# ── ข้อมูลแปลก ๆ ต้องได้ 400 ไม่ใช่ 500 ─────────────────────────────
class WeirdInputTests(AdvBase):
    def test_weird_values_never_500(self):
        cases = {
            'attendees float': {'attendees': 10.5},
            'attendees bool': {'attendees': True},
            'attendees huge': {'attendees': 10 ** 30},
            'attendees list': {'attendees': [10]},
            'title null byte': {'title': 'abc\x00def'},
            'title whitespace': {'title': '   '},
            'title list': {'title': ['a', 'b']},
            'title object': {'title': {'$ne': 1}},
            'room string': {'room': 'abc'},
            'room negative': {'room': -1},
            'year 9999 end': {'end_time': '9999-12-31T23:59:59Z'},
            'year 1 start': {'start_time': '0001-01-01T00:00:00Z'},
            'no timezone garbage': {'start_time': '2026-13-45T25:61:00'},
            'sql in title': {'title': "x'); DROP TABLE booking_booking;--"},
        }
        for name, over in cases.items():
            with self.subTest(name):
                r = self.book(**over)
                self.assertLess(r.status_code, 500, f'{name}: {r.status_code}')
        self.assertTrue(Booking.objects.model.objects.exists() or True)

    def test_sql_text_is_stored_verbatim(self):
        r = self.book(title="x'); DROP TABLE booking_booking;--")
        self.assertEqual(r.status_code, 201)
        self.assertEqual(Booking.objects.get().title, "x'); DROP TABLE booking_booking;--")

    def test_search_weird_values_never_500(self):
        for payload in (
            {'attendees': 10 ** 30, 'date': str(self.tue), 'start_time': '13:00', 'end_time': '15:00'},
            {'attendees': 10, 'date': '9999-12-31', 'start_time': '13:00', 'end_time': '15:00'},
            {'attendees': 10, 'date': str(self.tue), 'start_time': '13:00', 'end_time': '15:00',
             'building_code': "' OR 1=1 --", 'room_type': '<script>'},
        ):
            with self.subTest(payload=payload):
                r = self.client.post('/api/rooms/search/', payload, format='json')
                self.assertLess(r.status_code, 500)

    def test_query_string_weirdness_never_500(self):
        for url in ('/api/rooms/suggestions/?q=%00', '/api/rooms/suggestions/?q=' + 'ก' * 5000,
                    '/api/bookings/?page=abc', '/api/bookings/?page=999999',
                    '/api/rooms/?building=%27%3B--', '/api/notifications/?page=-1'):
            with self.subTest(url=url):
                self.assertLess(self.client.get(url).status_code, 500)

    def test_booking_forever_is_rejected(self):
        """จองยาวหลายปีทำให้ห้องถูกกันไว้ตลอด (สถานะรออนุมัติก็กันช่วงเวลาแล้ว)"""
        r = self.book(end_time=aware(self.tue + timedelta(days=365 * 5), 15).isoformat())
        self.assertEqual(r.status_code, 400, 'จองห้องยาว 5 ปีได้')


class DurationLimitTests(AdvBase):
    def test_booking_at_limit_ok_one_minute_over_rejected(self):
        from django.conf import settings
        start = aware(self.tue, 8)
        limit = timedelta(days=settings.MAX_BOOKING_DAYS)
        self.assertEqual(self.book(start_time=start.isoformat(), end_time=(start + limit).isoformat()).status_code, 201)
        Booking.objects.all().delete()
        r = self.book(start_time=start.isoformat(), end_time=(start + limit + timedelta(minutes=1)).isoformat())
        self.assertEqual(r.status_code, 400)

    def test_edit_cannot_stretch_past_limit(self):
        r = self.book()
        b = Booking.objects.get()
        far = aware(self.tue + timedelta(days=400), 15).isoformat()
        r = self.client.patch(f'/api/bookings/{b.id}/', {'end_time': far}, format='json')
        self.assertEqual(r.status_code, 400)

    def test_term_booking_over_a_year_rejected(self):
        r = self.client.post('/api/term-bookings/', {
            'room': self.room.id, 'subject_name': 'วิชายาวนาน', 'attendees': 20, 'day_of_week': 1,
            'start_time': '10:00', 'end_time': '12:00',
            'term_start': str(self.tue), 'term_end': str(self.tue + timedelta(days=400)),
        }, format='json')
        self.assertEqual(r.status_code, 400)


# ── token และการยืนยันตัวตน ────────────────────────────────────────
class TokenAttackTests(AdvBase):
    def test_garbage_and_tampered_tokens(self):
        good = str(RefreshToken.for_user(self.bot).access_token)
        head, body, sig = good.split('.')
        for name, token in {
            'garbage': 'abc.def.ghi',
            'alg none': f'{head}.{body}.',
            'bad signature': f'{head}.{body}.{sig[::-1]}',
        }.items():
            with self.subTest(name):
                self.assertEqual(self.bearer(token).get('/api/auth/profile/').status_code, 401)

    def test_refresh_token_cannot_be_used_as_access(self):
        refresh = str(RefreshToken.for_user(self.bot))
        self.assertEqual(self.bearer(refresh).get('/api/auth/profile/').status_code, 401)

    def test_deleted_or_disabled_user_token_rejected(self):
        token = str(RefreshToken.for_user(self.bot).access_token)
        self.bot.is_active = False
        self.bot.save()
        self.assertEqual(self.bearer(token).get('/api/auth/profile/').status_code, 401)

    def test_disabled_user_cannot_log_in(self):
        self.bot.is_active = False
        self.bot.save()
        self.client.force_authenticate(user=None)
        r = self.client.post('/api/auth/login/', {'username': 'qa_test_bot', 'password': 'QaTest1234'},
                             format='json')
        self.assertNotEqual(r.status_code, 200, 'บัญชีที่ถูกระงับยังล็อกอินได้')

    def test_old_tokens_die_after_password_change(self):
        """ขโมย token ไปแล้ว เจ้าของเปลี่ยนรหัสผ่าน — token เก่าต้องใช้ไม่ได้"""
        stolen = str(RefreshToken.for_user(self.bot).access_token)
        self.client.post('/api/auth/change-password/', {
            'old_password': 'QaTest1234', 'new_password': 'NewPass123', 'new_password2': 'NewPass123',
        }, format='json')
        self.assertEqual(self.bearer(stolen).get('/api/auth/profile/').status_code, 401,
                         'token ที่ออกก่อนเปลี่ยนรหัสผ่านยังใช้ได้')

    def test_register_mass_assignment_ignored(self):
        self.client.force_authenticate(user=None)
        self.client.post('/api/auth/register/', {
            'username': 'qa_sneaky', 'email': 'qa_sneaky@example.com', 'password': 'QaTest1234',
            'password2': 'QaTest1234', 'role': 'student', 'is_staff': True, 'is_superuser': True,
        }, format='json')
        u = User.objects.get(username='qa_sneaky')
        self.assertFalse(u.is_staff or u.is_superuser)

    def test_register_lookalike_usernames_rejected(self):
        """ชื่อที่หน้าตาเหมือนกันแต่ตัวพิมพ์/ตัวอักษรเต็มความกว้างต่างกัน ใช้ปลอมตัวได้"""
        self.client.force_authenticate(user=None)
        for name in ('QA_Test_Bot', 'ｑａ_test_bot'):
            with self.subTest(name=name):
                r = self.client.post('/api/auth/register/', {
                    'username': name, 'email': 'x@example.com', 'password': 'QaTest1234',
                    'password2': 'QaTest1234', 'role': 'student',
                }, format='json')
                self.assertEqual(r.status_code, 400, f'สมัคร {name!r} ได้ทั้งที่มี qa_test_bot อยู่แล้ว')

    def test_avatar_must_be_an_image(self):
        fake = SimpleUploadedFile('evil.png', b'<svg onload=alert(1)>', content_type='image/png')
        r = self.client.patch('/api/auth/profile/', {'avatar': fake}, format='multipart')
        self.assertEqual(r.status_code, 400)


# ── การยกเลิก/เช็คอินในสถานะที่ไม่ควรทำได้ ─────────────────────────
class StateAbuseTests(AdvBase):
    def test_cannot_cancel_finished_or_rejected(self):
        for status in ('completed', 'rejected'):
            with self.subTest(status=status):
                b = Booking.objects.create(user=self.bot, room=self.room, title='x', attendees=5,
                                           start_time=aware(self.tue, 8), end_time=aware(self.tue, 9),
                                           status=status)
                r = self.client.post(f'/api/bookings/{b.id}/cancel/')
                b.refresh_from_db()
                self.assertEqual(r.status_code, 400, f'ยกเลิกรายการ {status} ได้')
                self.assertEqual(b.status, status)

    def test_email_checkin_link_needs_a_click(self):
        """โปรแกรมสแกนลิงก์ในอีเมล (เปิด GET ให้อัตโนมัติ) ต้องไม่ทำให้เช็คอินเอง"""
        start = timezone.now() + timedelta(minutes=10)
        b = Booking.objects.create(user=self.bot, room=self.room, title='x', attendees=5,
                                   start_time=start, end_time=start + timedelta(hours=1),
                                   status='approved')
        anon = APIClient()
        r = anon.get(f'/api/bookings/{b.id}/checkin/{b.checkin_token}/')
        b.refresh_from_db()
        self.assertEqual(r.status_code, 200)
        self.assertEqual(b.status, 'approved', 'แค่เปิดลิงก์ (GET) ก็เช็คอินแล้ว')
        r = anon.post(f'/api/bookings/{b.id}/checkin/{b.checkin_token}/')
        b.refresh_from_db()
        self.assertEqual(r.status_code, 200)
        self.assertEqual(b.status, 'checked_in')


# ── export Excel ────────────────────────────────────────────────────
class ExportLeakTests(AdvBase):
    def export(self):
        from openpyxl import load_workbook
        self.client.force_authenticate(user=self.admin)
        r = self.client.get('/api/export/excel/?sheets=users,bookings')
        self.assertEqual(r.status_code, 200)
        return load_workbook(io.BytesIO(r.content))

    def test_export_has_no_password_hashes_or_tokens(self):
        wb = self.export()
        headers = {c.value for ws in wb.worksheets for c in ws[1]}
        self.assertNotIn('password', headers, 'export มี hash รหัสผ่านของผู้ใช้ทุกคน')
        self.assertNotIn('checkin_token', headers, 'export มี token ลิงก์เช็คอิน/ยกเลิกของทุกการจอง')

    def test_export_does_not_turn_titles_into_formulas(self):
        self.book(title='=HYPERLINK("http://evil.example","คลิก")')
        wb = self.export()
        cells = [c for ws in wb.worksheets for row in ws.iter_rows() for c in row
                 if isinstance(c.value, str) and 'HYPERLINK' in c.value]
        self.assertTrue(cells)
        for c in cells:
            self.assertNotEqual(c.data_type, 'f', 'ชื่อกิจกรรมกลายเป็นสูตร Excel')


# ── ยึดบัญชีด้วยรหัสนักศึกษาคนอื่น ─────────────────────────────────
class AccountClaimTests(AdvBase):
    def test_pre_registered_student_id_does_not_survive_ldap_claim(self):
        """คนร้ายสมัคร username = รหัสนักศึกษาคนอื่นไว้ก่อน เก็บ token ไว้ พอเจ้าของ
        ตัวจริงล็อกอินผ่าน LDAP (ระบบรวมเป็นบัญชีเดียวกัน) token ของคนร้ายต้องใช้ไม่ได้"""
        from unittest import mock
        self.client.force_authenticate(user=None)
        # สมัครแบบนี้ผ่าน API ไม่ได้แล้ว (test_register_student_id_username_rejected)
        # แต่บัญชีเก่าที่สมัครไว้ก่อนแก้ยังมีอยู่ได้ — จำลองด้วยการสร้างตรง
        User.objects.create_user(username='6699999999', password='Attack1234',
                                 email='attacker@example.com', role='student')
        login = self.client.post('/api/auth/login/', {'username': '6699999999', 'password': 'Attack1234'},
                                 format='json')
        attacker_access = login.data['access']

        ldap_ok = {'username': '6699999999', 'full_name': 'เจ้าของ ตัวจริง',
                   'email': '6699999999@ubu.ac.th', 'department': 'วิศวกรรมศาสตร์'}
        with mock.patch('booking.serializers.authenticate_ldap', return_value=ldap_ok):
            victim = self.client.post('/api/auth/login/', {'username': '6699999999', 'password': 'real-pass'},
                                      format='json')
        self.assertEqual(victim.status_code, 200)
        self.assertEqual(self.bearer(attacker_access).get('/api/auth/profile/').status_code, 401,
                         'คนร้ายยังใช้ token เดิมเข้าบัญชีของเจ้าของตัวจริงได้')


class WebSocketAuthTests(TransactionTestCase):
    """WebSocket แจ้งเตือนใช้กฎ token เดียวกับ REST API"""

    def setUp(self):
        self.user = User.objects.create_user(username='qa_test_bot', password='QaTest1234', role='student')

    def connects(self, token):
        from asgiref.sync import async_to_sync
        from channels.testing import WebsocketCommunicator
        from room_booking.asgi import application

        async def attempt():
            comm = WebsocketCommunicator(application, f'/ws/notifications/?token={token}')
            ok, _ = await comm.connect()
            await comm.disconnect()
            return ok
        return async_to_sync(attempt)()

    def test_valid_token_connects(self):
        from booking.authentication import issue_tokens
        self.assertTrue(self.connects(str(issue_tokens(self.user).access_token)))

    def test_stale_or_disabled_token_rejected(self):
        from booking.authentication import issue_tokens
        token = str(issue_tokens(self.user).access_token)
        self.user.set_password('Changed123')
        self.user.save()
        self.assertFalse(self.connects(token), 'token เก่าหลังเปลี่ยนรหัสยังเชื่อม WebSocket ได้')
        token = str(issue_tokens(self.user).access_token)
        self.user.is_active = False
        self.user.save()
        self.assertFalse(self.connects(token), 'บัญชีที่ถูกระงับยังเชื่อม WebSocket ได้')

    def test_no_token_rejected(self):
        self.assertFalse(self.connects(''))
